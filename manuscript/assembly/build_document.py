"""Build the public working-manuscript DOCX from the manuscript's Markdown source.

Requires python-docx. Rendering is a separate verification step.
"""
from pathlib import Path
import argparse
import hashlib
import json
import re
from datetime import datetime, timezone

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "spea_comparative_manuscript.md"
OUTPUT = ROOT / "rebuilt/Spea_Comparative_Manuscript.docx"
RUN = ROOT / "rebuilt"


def xml(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def hyperlink(p, label, target):
    relation = p.part.relate_to(target, RT.HYPERLINK, is_external=True)
    link = xml("w:hyperlink", **{"r:id": relation})
    run = xml("w:r")
    properties = xml("w:rPr")
    properties.append(xml("w:color", **{"w:val": "000000"}))
    run.append(properties)
    value = xml("w:t")
    value.text = label
    run.append(value)
    link.append(run)
    p._p.append(link)


def inline(p, text):
    pattern = r"\*\*(.+?)\*\*|\*(.+?)\*|\[([^\]]+)\]\(([^)]+)\)"
    end = 0
    for m in re.finditer(pattern, text):
        if m.start() > end:
            p.add_run(text[end:m.start()])
        if m.group(1) is not None:
            p.add_run(m.group(1)).bold = True
        elif m.group(2) is not None:
            p.add_run(m.group(2)).italic = True
        else:
            hyperlink(p, m.group(3), m.group(4))
        end = m.end()
    if end < len(text):
        p.add_run(text[end:])


def mr(text):
    e = xml("m:r")
    t = xml("m:t")
    t.text = text
    e.append(t)
    return e


def sub(base, index):
    node = xml("m:sSub")
    e, s = xml("m:e"), xml("m:sub")
    e.append(mr(base))
    s.append(mr(index))
    node.extend([e, s])
    return node


def density_equation(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    para, math = xml("m:oMathPara"), xml("m:oMath")
    math.extend([sub("ℓ", "i"), mr(" = "), sub("η", "i"), mr(" − log(")])
    fraction, numerator, denominator = xml("m:f"), xml("m:num"), xml("m:den")
    numerator.extend([sub("∑", "j"), sub("a", "j"), mr(" exp("), sub("η", "j"), mr(")")])
    denominator.extend([sub("∑", "j"), sub("a", "j")])
    fraction.extend([numerator, denominator])
    math.extend([fraction, mr(")")])
    para.append(math)
    p._p.append(para)


def table(doc, rows):
    widths = {3: [1.15, 2.7, 2.65], 6: [1.8, .9, .9, .6, .9, 1.4], 4: [2, 1.5, 1.5, 1.5]}[len(rows[0])]
    t = doc.add_table(rows=0, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    props = t._tbl.tblPr
    borders = xml("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        borders.append(xml("w:" + side, **{"w:val": "single", "w:sz": "4", "w:color": "D9D9D9"}))
    props.append(borders)
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for i, values in enumerate(rows):
        row = t.add_row()
        row._tr.get_or_add_trPr().append(xml("w:cantSplit"))
        if i == 0:
            row._tr.get_or_add_trPr().append(xml("w:tblHeader"))
        for j, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            margins = xml("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                margins.append(xml("w:" + side, **{"w:w": "85", "w:type": "dxa"}))
            pr.append(margins)
            if i == 0:
                pr.append(xml("w:shd", **{"w:fill": "EAEFF2"}))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.08
            p.paragraph_format.keep_with_next = i < len(rows) - 1
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 or len(widths) == 3 else WD_ALIGN_PARAGRAPH.CENTER
            inline(p, value)
            for run in p.runs:
                run.font.size = Pt(10)
                if i == 0:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--receipt", type=Path, default=RUN / "document_build.json")
    args = parser.parse_args()
    source, output = args.source.resolve(), args.output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(.8), Inches(.75)
    sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(.32)
    for style in doc.styles:
        if style.type == 1:
            style.font.name = "Times New Roman"
            style.font.color.rgb = RGBColor(0, 0, 0)
            for node in style.element.xpath(".//w:rFonts"):
                for key in list(node.attrib):
                    if key.endswith("Theme"):
                        del node.attrib[key]
            for node in style.element.xpath(".//w:pBdr"):
                node.getparent().remove(node)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True
    for name, size, before, after in [("Title", 20, 0, 12), ("Heading 1", 14, 16, 7), ("Heading 2", 12, 12, 5)]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    doc.styles["Caption"].font.name = "Times New Roman"
    doc.styles["Caption"].font.size = Pt(10.5)
    doc.styles["Caption"].font.italic = False
    doc.styles["Caption"].font.bold = False
    doc.styles["Caption"].paragraph_format.line_spacing = 1.08
    header = sec.header.paragraphs[0]
    header.text = "Climatic differences and geographic prediction across Spea"
    header.runs[0].font.size = Pt(9)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Working manuscript  |  ").font.size = Pt(9)
    field = xml("w:fldSimple", **{"w:instr": "PAGE"})
    footer._p.append(field)
    sec.different_first_page_header_footer = True
    doc.settings.element.append(xml("w:updateFields", **{"w:val": "true"}))
    doc.core_properties.title = "Recorded climatic differences and geographic prediction across Spea"
    doc.core_properties.author = "Kevin M. Neal"
    doc.core_properties.subject = "Working manuscript using completed comparative analyses"
    # Release date metadata, fixed for reproducible DOCX member payloads.
    doc.core_properties.created = datetime(2026, 10, 29, tzinfo=timezone.utc)
    doc.core_properties.modified = datetime(2026, 10, 29, tzinfo=timezone.utc)
    text = source.read_text().replace("–", "-").replace("—", "-").replace("\u2011", "-")
    lines = text.splitlines()
    i = 0
    references = False
    supporting = False
    figures = []
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line == "<!-- PAGEBREAK -->":
            doc.add_page_break(); i += 1; continue
        if line == "$$":
            i += 1
            while i < len(lines) and lines[i].strip() != "$$": i += 1
            density_equation(doc); i += 1; continue
        image = re.fullmatch(r"!\[.*\]\((.*)\)", line)
        if image:
            path = source.parent / image.group(1)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(6)
            picture = p.add_run().add_picture(str(path), width=Inches(6.5))
            picture._inline.docPr.set("descr", re.match(r"!\[(.*)\]", line).group(1))
            figures.append(str(path.relative_to(source.parent)))
            i += 1; continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                values = [v.strip() for v in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", v) for v in values): rows.append(values)
                i += 1
            table(doc, rows); continue
        heading = re.match(r"^(#{1,3})\s+(.*)", line)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2)
            p = doc.add_paragraph(style="Title" if level == 1 else f"Heading {level-1}")
            inline(p, title)
            references = title == "References"
            if title == "Supporting information":
                supporting = True
            if references: p.paragraph_format.page_break_before = True
            i += 1; continue
        text_lines = [line]
        i += 1
        while i < len(lines) and lines[i].strip():
            text_lines.append(lines[i].strip()); i += 1
        content = " ".join(text_lines)
        is_caption = content.startswith("**Figure ") or content.startswith("**Table ")
        p = doc.add_paragraph(style="Caption" if is_caption else "Normal")
        inline(p, content)
        if supporting and not is_caption:
            p.paragraph_format.line_spacing = 1.07
            for run in p.runs:
                run.font.size = Pt(11)
        if content.startswith("**Table "):
            p.paragraph_format.keep_with_next = True
        if references:
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.03
            p.paragraph_format.keep_together = True
            p.paragraph_format.left_indent = Inches(.22)
            p.paragraph_format.first_line_indent = Inches(-.22)
            for run in p.runs: run.font.size = Pt(10)
    doc.save(output)
    sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
    receipt = {"source": str(source), "source_sha256": sha(source),
               "output": str(output), "output_sha256": sha(output),
               "paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "figures": figures,
               "native_word_equations": 1, "format": "Letter portrait, 11.5 point body, 11 point supporting text, review figures after references",
               "source_words": len(text.split()), "render_status": "not_yet_rendered"}
    args.receipt.parent.mkdir(parents=True, exist_ok=True)
    args.receipt.write_text(json.dumps(receipt, indent=2) + "\n")
    print(json.dumps(receipt, indent=2))


if __name__ == "__main__":
    main()
